"""Geradores de arquivos reais para os testes (sem fixtures binárias no repo)."""

import io

import docx


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(
    pages: list[list[str]], info: dict[str, str] | None = None, xmp: str | None = None
) -> bytes:
    """PDF mínimo com Helvetica/WinAnsi, uma linha de texto por item.

    `info` vira o dicionário /Info do trailer; `xmp`, o stream /Metadata do catálogo.
    """

    def esc(line: str) -> str:
        out = []
        for ch in line.encode("cp1252"):
            c = chr(ch)
            if c in "()\\":
                out.append("\\" + c)
            elif ch > 126:
                out.append(f"\\{ch:03o}")
            else:
                out.append(c)
        return "".join(out)

    n = len(pages)
    page_ids = [4 + 2 * i for i in range(n)]
    next_id = 4 + 2 * n
    info_id = xmp_id = None
    if info:
        info_id, next_id = next_id, next_id + 1
    if xmp:
        xmp_id = next_id

    kids = " ".join(f"{p} 0 R" for p in page_ids)
    catalog = "<< /Type /Catalog /Pages 2 0 R"
    catalog += f" /Metadata {xmp_id} 0 R >>" if xmp_id else " >>"
    objects: dict[int, bytes] = {
        1: catalog.encode(),
        2: f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    for pid, lines in zip(page_ids, pages, strict=True):
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td"]
        for line in lines:
            ops.append(f"({esc(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects[pid] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {pid + 1} 0 R >>"
        ).encode()
        objects[pid + 1] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    if info_id:
        entries = " ".join(f"/{k} ({esc(v)})" for k, v in info.items())
        objects[info_id] = f"<< {entries} >>".encode()
    if xmp_id:
        body = xmp.encode()
        objects[xmp_id] = (
            f"<< /Type /Metadata /Subtype /XML /Length {len(body)} >>\nstream\n".encode()
            + body
            + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for oid in sorted(objects):
        offsets[oid] = len(out)
        out += f"{oid} 0 obj\n".encode() + objects[oid] + b"\nendobj\n"
    xref = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for oid in range(1, size):
        out += f"{offsets[oid]:010d} 00000 n \n".encode()
    trailer = f"<< /Size {size} /Root 1 0 R"
    trailer += f" /Info {info_id} 0 R >>" if info_id else " >>"
    out += f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)
