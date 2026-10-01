import io
import zipfile

import docx
import pytest
from docx.oxml import parse_xml
from pypdf import PdfReader

from app.services.extract import ExtractionError, FileKind, extract
from app.services.sanitize import sanitize_file, sanitize_text
from tests.factories import make_docx, make_pdf

# --- texto ----------------------------------------------------------------------


def test_removes_invisible_chars():
    text, changes = sanitize_text("pa​la‌vra﻿ com­binada⁠")
    assert text == "palavra combinada"
    assert changes.invisible_chars == 5


def test_zwj_kept_inside_emoji_removed_between_letters():
    family = "\U0001f468‍\U0001f469"
    text, changes = sanitize_text(f"te‍xto {family}")
    assert text == f"texto {family}"
    assert changes.invisible_chars == 1


def test_normalizes_typography():
    text, changes = sanitize_text("“Citação” e ‘outra’ — fim… 1990–2000, ida—volta")
    assert text == "\"Citação\" e 'outra' - fim... 1990-2000, ida - volta"
    assert changes.quotes == 4
    assert changes.dashes == 3
    assert changes.ellipses == 1
    assert changes.special_spaces == 1
    assert changes.total == 9


def test_clean_text_is_untouched():
    text = 'Texto comum, sem nada - nem aspas "curvas".'
    assert sanitize_text(text) == (text, sanitize_text("")[1])


# --- txt ------------------------------------------------------------------------


def test_txt_is_returned_as_utf8():
    result = sanitize_file("Olá “mundo”​".encode("utf-16"), "a.txt")
    assert result.kind is FileKind.TXT
    assert result.data.decode("utf-8") == 'Olá "mundo"'
    assert result.report.text.total == 3


# --- docx -----------------------------------------------------------------------

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def docx_with_traces() -> bytes:
    document = docx.Document(io.BytesIO(make_docx("Texto “colado”​ do chat—bot…")))
    document.core_properties.author = "Fulano de Tal"
    document.core_properties.last_modified_by = "Fulano de Tal"
    document.core_properties.title = "Meu TCC"
    paragraph = document.add_paragraph()
    run = paragraph.add_run("revisado")
    document.add_comment(run, text="Ver isso", author="Fulano de Tal", initials="FT")
    # Alteração controlada (w:ins) com autor.
    paragraph._p.append(
        parse_xml(
            f'<w:ins xmlns:w="{W}" w:id="99" w:author="Fulano de Tal" '
            f'w:date="2026-09-30T10:00:00Z"><w:r><w:t> inserido</w:t></w:r></w:ins>'
        )
    )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_cleans_text_metadata_and_authors():
    result = sanitize_file(docx_with_traces(), "tcc.docx")
    report = result.report

    assert result.kind is FileKind.DOCX
    assert extract(result.data, "tcc.docx").text.startswith(
        'Texto "colado" do chat - bot...\n\nrevisado inserido'
    )
    assert report.text.invisible_chars == 1
    assert {"creator", "lastModifiedBy", "created", "modified", "revision", "description"} <= set(
        report.metadata_removed
    )
    assert {"Template", "TotalTime"} <= set(report.metadata_removed)
    assert report.authors_anonymized >= 3  # comentário (author + initials) + w:ins

    with zipfile.ZipFile(io.BytesIO(result.data)) as zf:
        assert all(i.date_time == (1980, 1, 1, 0, 0, 0) for i in zf.infolist())
        everything = b"".join(zf.read(n) for n in zf.namelist())
    assert b"Fulano" not in everything
    assert b"python-docx" not in everything

    # Continua um .docx válido, e o que é conteúdo do usuário fica.
    reopened = docx.Document(io.BytesIO(result.data))
    assert reopened.core_properties.author == ""
    assert reopened.core_properties.title == "Meu TCC"


def test_docx_is_idempotent():
    once = sanitize_file(docx_with_traces(), "a.docx")
    twice = sanitize_file(once.data, "a.docx")
    assert twice.report.text.total == 0
    assert twice.report.authors_anonymized == 0
    assert twice.report.metadata_removed == []


def test_docx_rejects_xml_entities():
    """XXE: entidade externa no XML não pode ser resolvida."""
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(make_docx("oi"))) as src, zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            content = src.read(info)
            if info.filename == "word/document.xml":
                content = content.replace(
                    b"?>",
                    b'?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>',
                    1,
                ).replace(b">oi<", b">&x;<")
            dst.writestr(info, content)
    result = sanitize_file(buf.getvalue(), "a.docx")
    assert b"root:" not in result.data


# --- pdf ------------------------------------------------------------------------

XMP = '<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF/></x:xmpmeta>'


def test_pdf_removes_info_and_xmp():
    pdf = make_pdf(
        [["Conteúdo da página."]],
        info={"Author": "Fulano", "Creator": "Microsoft Word", "Producer": "Gerador X"},
        xmp=XMP,
    )
    result = sanitize_file(pdf, "a.pdf")

    assert set(result.report.metadata_removed) == {"Author", "Creator", "Producer", "XMP"}
    assert result.report.notes  # avisa que o texto não foi tocado
    reader = PdfReader(io.BytesIO(result.data))
    assert not reader.metadata
    assert "/Metadata" not in reader.trailer["/Root"]
    for marker in (b"Fulano", b"Microsoft Word", b"Gerador X", b"xmpmeta"):
        assert marker not in result.data
    assert extract(result.data, "a.pdf").text == "Conteúdo da página."


def test_pdf_with_password_is_rejected():
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf([["x"]]))))
    writer.encrypt("segredo")
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ExtractionError, match="senha"):
        sanitize_file(buf.getvalue(), "a.pdf")
