import io
import zipfile

import docx
import pytest

from app.services.extract import (
    ExtractionError,
    FileKind,
    count_words,
    detect_kind,
    extract,
)
from tests.factories import make_docx, make_pdf

# --- contagem -------------------------------------------------------------------


def test_count_words_treats_compounds_as_one():
    assert count_words("O guarda-chuva d'água, 2026!") == 4
    assert count_words("") == 0


# --- txt ------------------------------------------------------------------------


def test_txt_utf8_with_bom_and_crlf():
    data = "﻿Primeiro parágrafo.\r\n\r\n\r\n\r\nSegundo.  \r\n".encode()
    result = extract(data, "texto.txt")
    assert result.kind is FileKind.TXT
    assert result.text == "Primeiro parágrafo.\n\nSegundo."
    assert result.word_count == 3


def test_txt_cp1252_fallback():
    data = "Ação “citação” — fim".encode("cp1252")
    assert extract(data, "notas.txt").text == "Ação “citação” — fim"


def test_txt_utf16():
    data = "Olá mundo".encode("utf-16")
    assert extract(data, "x.txt").text == "Olá mundo"


def test_txt_keeps_zero_width_for_sanitize():
    assert "​" in extract("pala​vra".encode(), "x.txt").text


def test_binary_with_txt_extension_is_rejected():
    with pytest.raises(ExtractionError):
        extract(b"\x00\x01\x02abc", "x.txt")


# --- docx -----------------------------------------------------------------------


def test_docx_paragraphs_and_table_in_order():
    data = make_docx("Introdução.", "", "Corpo do texto.", table=[["A", "B"], ["C", "D"]])
    result = extract(data, "trabalho.docx")
    assert result.kind is FileKind.DOCX
    assert result.mime_type.endswith("wordprocessingml.document")
    assert result.text == "Introdução.\n\nCorpo do texto.\n\nA\n\nB\n\nC\n\nD"


def test_docx_detected_by_content_not_extension():
    assert detect_kind(make_docx("oi"), "renomeado.pdf") is FileKind.DOCX


def test_empty_docx_is_rejected():
    with pytest.raises(ExtractionError, match="Não encontramos texto"):
        extract(make_docx(), "vazio.docx")


def test_docx_zip_bomb_is_rejected(monkeypatch):
    monkeypatch.setattr("app.services.extract.MAX_DOCX_UNCOMPRESSED_BYTES", 1000)
    with pytest.raises(ExtractionError, match="grande demais"):
        extract(make_docx("x" * 5000), "a.docx")


def test_other_zip_formats_are_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("content.xml", "<x/>")
    with pytest.raises(ExtractionError, match=".odt"):
        detect_kind(buf.getvalue(), "a.odt")


def test_legacy_doc_and_rtf_are_rejected():
    with pytest.raises(ExtractionError, match=".doc antigo"):
        detect_kind(b"\xd0\xcf\x11\xe0" + b"\x00" * 100, "a.doc")
    with pytest.raises(ExtractionError, match=".rtf"):
        detect_kind(b"{\\rtf1 ola}", "a.rtf")


def test_unknown_extension_is_rejected():
    with pytest.raises(ExtractionError, match="Formato não suportado"):
        detect_kind(b"qualquer coisa", "imagem.png")


# --- pdf ------------------------------------------------------------------------

FULL = "Este é um parágrafo longo que ocupa a largura inteira da linha do"


def test_pdf_rebuilds_paragraphs_and_hyphenation():
    pdf = make_pdf(
        [
            [
                FULL,
                "documento e continua na linha de baixo sem quebrar o parágrafo em",
                "dois pedaços, até terminar aqui.",
                "Segundo parágrafo começa nesta linha e também é bem compri-",
                "do, terminando na linha seguinte.",
                "1",
            ],
            ["Terceiro, na outra página."],
        ]
    )
    result = extract(pdf, "x.pdf")
    assert result.kind is FileKind.PDF
    assert result.page_count == 2
    paragraphs = result.text.split("\n\n")
    assert paragraphs == [
        f"{FULL} documento e continua na linha de baixo sem quebrar o parágrafo em "
        "dois pedaços, até terminar aqui.",
        "Segundo parágrafo começa nesta linha e também é bem comprido, "
        "terminando na linha seguinte.",
        "Terceiro, na outra página.",
    ]


def test_pdf_without_text_is_rejected():
    with pytest.raises(ExtractionError, match="digitalizado"):
        extract(make_pdf([[]]), "scan.pdf")


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ExtractionError, match="corrompido"):
        extract(b"%PDF-1.4\nlixo sem estrutura", "x.pdf")


def test_upload_size_limit(monkeypatch):
    monkeypatch.setattr("app.services.extract.settings.max_upload_mb", 0)
    with pytest.raises(ExtractionError, match="limite"):
        extract(b"abc", "x.txt")


def test_docx_merged_cell_is_not_duplicated():
    document = docx.Document()
    t = document.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Mesclada"
    t.cell(1, 0).text, t.cell(1, 1).text = "X", "Y"
    buf = io.BytesIO()
    document.save(buf)
    assert extract(buf.getvalue(), "t.docx").text == "Mesclada\n\nX\n\nY"


def test_docx_tracked_changes_follow_what_word_shows():
    from docx.oxml import parse_xml

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    document = docx.Document(io.BytesIO(make_docx("Antes")))
    p = document.paragraphs[0]._p
    p.append(
        parse_xml(f'<w:ins xmlns:w="{w}" w:id="1" w:author="a"><w:r><w:t> novo</w:t></w:r></w:ins>')
    )
    p.append(
        parse_xml(
            f'<w:del xmlns:w="{w}" w:id="2" w:author="a">'
            "<w:r><w:delText> velho</w:delText></w:r></w:del>"
        )
    )
    p.append(parse_xml(f'<w:r xmlns:w="{w}"><w:tab/><w:t>fim</w:t></w:r>'))
    buf = io.BytesIO()
    document.save(buf)
    assert extract(buf.getvalue(), "a.docx").text == "Antes novo\tfim"
